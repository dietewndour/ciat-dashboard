"""Exports Excel et PDF."""
import io
from datetime import datetime
from xml.sax.saxutils import escape as esc

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

import analytics as an

NAVY = "#12355B"


def _filters_text(db, f):
    names = []
    if "cohort_id" in f:
        names.append("cohorte : " + db.execute("SELECT name FROM cohorts WHERE id=?", (f["cohort_id"],)).fetchone()[0])
    if "domain_id" in f:
        names.append("domaine : " + db.execute("SELECT name FROM domains WHERE id=?", (f["domain_id"],)).fetchone()[0])
    if "status" in f: names.append("statut : " + an.STATUS_LABELS[f["status"]])
    if "gender" in f: names.append("genre : " + an.GENDER_LABELS[f["gender"]])
    if "date_from" in f: names.append("dépôt depuis le " + f["date_from"])
    if "date_to" in f: names.append("dépôt jusqu'au " + f["date_to"])
    if "owner" in f: names.append("porteur : " + f["owner"])
    return ", ".join(names) or "aucun filtre"


def excel(db, f):
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY[1:])

    def table(ws, header, rows, start=1):
        for j, h in enumerate(header, 1):
            c = ws.cell(row=start, column=j, value=h); c.font = head; c.fill = fill
        for i, r in enumerate(rows, start + 1):
            for j, v in enumerate(r, 1):
                ws.cell(row=i, column=j, value=v)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                60, max(12, max(len(str(c.value or "")) for c in col) + 2))

    ws = wb.active; ws.title = "Synthèse"
    ws["A1"] = "Tableau de bord du CIAT - synthèse"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Généré le {datetime.now():%d/%m/%Y %H:%M}. Filtres : {_filters_text(db, f)}"
    k = an.kpis(db, f)
    table(ws, ["Indicateur", "Valeur"], [(v["label"], an.fmt(v["value"], v["unit"])) for v in k.values()], 4)
    for name, group in (("Par domaine", "domain"), ("Par cohorte", "cohort"), ("Par genre", "gender")):
        w = wb.create_sheet(name)
        data = {m: dict(an.aggregate(db, f, m, group)) for m in ("projects", "selected", "prototypes", "delayed", "funding")}
        labels = list(data["projects"].keys())
        table(w, [an.GROUPS[group][0], "Idées", "Sélectionnés", "Prototypes", "En retard", "Financement (FCFA)"],
              [(l, *(data[m].get(l, 0) for m in ("projects", "selected", "prototypes", "delayed", "funding"))) for l in labels])
    w = wb.create_sheet("Projets")
    rows = an.list_projects(db, f, limit=5000)
    table(w, ["ID", "Titre", "Porteur", "Genre", "Domaine", "Cohorte", "Statut", "Déposé le",
              "Progression %", "Financement (FCFA)", "En retard"],
          [(r["id"], r["title"], r["owner_name"], r["owner_gender"], r["domain"], r["cohort"],
            an.STATUS_LABELS[r["status"]], r["submitted_on"], r["progress_pct"], r["funding"],
            "oui" if r["delayed"] else "non") for r in rows])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def pdf(db, f):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm, title="Synthèse CIAT")
    st = getSampleStyleSheet()
    story = [Paragraph("Tableau de bord du CIAT - synthèse", st["Title"]),
             Paragraph(esc(f"Généré le {datetime.now():%d/%m/%Y %H:%M}. Filtres : {_filters_text(db, f)}."), st["Normal"]),
             Paragraph("Données fictives de démonstration.", st["Italic"]), Spacer(1, 12)]

    def tbl(data, widths=None):
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F4F8")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        return t

    k = an.kpis(db, f)
    story += [Paragraph("Indicateurs clés", st["Heading2"]),
              tbl([["Indicateur", "Valeur"]] + [[v["label"], an.fmt(v["value"], v["unit"])] for v in k.values()],
                  [11 * cm, 5 * cm]), Spacer(1, 12)]
    for title, group in (("Répartition par domaine", "domain"), ("Comparaison des cohortes", "cohort")):
        d = {m: dict(an.aggregate(db, f, m, group)) for m in ("projects", "selected", "prototypes")}
        data = [[an.GROUPS[group][0], "Idées", "Sélect.", "Prototypes"]]
        for label in d["projects"]:
            data.append([Paragraph(esc(label), st["BodyText"]), d["projects"].get(label, 0),
                         d["selected"].get(label, 0), d["prototypes"].get(label, 0)])
        story += [Paragraph(title, st["Heading2"]), tbl(data, [8.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm]), Spacer(1, 12)]
    late = an.delayed_projects(db, f, 15)
    story.append(Paragraph("Projets en retard", st["Heading2"]))
    if late:
        story.append(tbl([["Projet", "Porteur", "Jalons en retard", "Plus ancien"]] +
                         [[Paragraph(esc(r["title"]), st["BodyText"]), r["owner_name"], r["late_count"], r["oldest_due"]] for r in late],
                         [6.5 * cm, 4.5 * cm, 3 * cm, 3 * cm]))
    else:
        story.append(Paragraph("Aucun projet en retard pour ce périmètre.", st["Normal"]))
    story += [Spacer(1, 10), Paragraph(
        "Limite : les retards se calculent sur les jalons non terminés dont l'échéance est passée. "
        "Un jalon absent de la base n'est pas détecté.", st["Italic"])]
    doc.build(story)
    return buf.getvalue()


def project_fiche(db, pid):
    p = db.execute("""SELECT p.*, d.name AS domain, c.name AS cohort FROM projects p
                      JOIN domains d ON d.id=p.domain_id JOIN cohorts c ON c.id=p.cohort_id WHERE p.id=?""", (pid,)).fetchone()
    if not p:
        return None
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm)
    st = getSampleStyleSheet()
    story = [Paragraph(esc(f"Fiche de synthèse : {p['title']}"), st["Title"]),
             Paragraph(f"Porteur : {esc(p['owner_name'])}<br/>Domaine : {p['domain']}<br/>Cohorte : {p['cohort']}<br/>"
                       f"Statut : {an.STATUS_LABELS[p['status']]}<br/>Déposé le : {p['submitted_on']}<br/>"
                       f"Progression : {p['progress_pct']} %", st["Normal"]), Spacer(1, 10)]
    ms = db.execute("SELECT label,due_on,done_on FROM milestones WHERE project_id=? ORDER BY due_on", (pid,)).fetchall()
    if ms:
        story.append(Paragraph("Jalons", st["Heading2"]))
        for m in ms:
            state = f"terminé le {m['done_on']}" if m["done_on"] else (
                "EN RETARD" if m["due_on"] < an.today().isoformat() else "à venir")
            story.append(Paragraph(f"- {esc(m['label'])} (échéance {m['due_on']}) : {state}", st["Normal"]))
    fu = db.execute("SELECT source,amount_fcfa,status FROM fundings WHERE project_id=?", (pid,)).fetchall()
    if fu:
        story.append(Paragraph("Financements", st["Heading2"]))
        for x in fu:
            story.append(Paragraph(f"- {esc(x['source'])} : {an.fmt(x['amount_fcfa'], 'fcfa')} ({x['status']})", st["Normal"]))
    pa = db.execute("SELECT partner_name,kind,status FROM partnerships WHERE project_id=?", (pid,)).fetchall()
    if pa:
        story.append(Paragraph("Partenariats", st["Heading2"]))
        for x in pa:
            story.append(Paragraph(f"- {esc(x['partner_name'])} ({x['kind']}, {x['status']})", st["Normal"]))
    doc.build(story)
    return buf.getvalue()
